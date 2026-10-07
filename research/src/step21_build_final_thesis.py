"""
step21_build_final_thesis.py
----------------------------
Combine the ten corrected chapter documents into THESIS_FINAL.docx.

The previous handover note said this could not be done programmatically without
damaging styles, figure numbering and headers, and left the combined draft
carrying every claim the chapter files had already corrected. docxcompose does
it properly: it merges numbering definitions, style ids and relationship ids
rather than concatenating XML, so images, captions and list numbering survive.

Each chapter is separated by a page break, which is how the chapters were laid
out when they were exported individually.

Verification after the merge, all of which must pass:
  * the combined file re-opens as a valid .docx
  * its image count equals the sum of the chapters' image counts
  * its table count equals the sum of the chapters' table counts
  * its body text contains every chapter's opening heading
  * none of the superseded values appears outside the passages that explicitly
    supersede them
"""

from __future__ import annotations

import os
import re
import sys
import zipfile

from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TF = os.path.join(ROOT, "thesis_final")
OUT = os.path.join(TF, "THESIS_FINAL.docx")

CHAPTERS = [
    "Chapter0_FrontMatter.docx",
    "Chapter1_Introduction.docx",
    "Chapter2_LiteratureReview.docx",
    "Chapter3_Methodology.docx",
    "Chapter4_SystemAnalysis.docx",
    "Chapter5_Implementation.docx",
    "Chapter6_TestingEvaluation.docx",
    "Chapter7_Conclusions.docx",
    "Chapter8_References.docx",
    "Chapter9_Appendices.docx",
]

# values that were corrected during the audit; each may appear only in a
# passage that names it as superseded
SUPERSEDED = {
    "41.25": ("superseded", "earlier", "previously", "300-tree"),
    "48.18": ("superseded", "earlier", "previously"),
    "6,211": ("silently", "superseded", "earlier"),
    "6211": ("mismatch", "silently", "superseded"),
}


def counts(path):
    z = zipfile.ZipFile(path)
    imgs = sum(1 for n in z.namelist() if n.startswith("word/media/"))
    z.close()
    d = Document(path)
    return imgs, len(d.tables)


def main():
    paths = [os.path.join(TF, c) for c in CHAPTERS]
    for p in paths:
        if not os.path.exists(p):
            raise SystemExit(f"missing chapter: {p}")

    want_imgs = want_tables = 0
    for p in paths:
        i, t = counts(p)
        want_imgs += i
        want_tables += t
        print(f"  {os.path.basename(p):34s} images {i:>2}  tables {t:>2}")

    master = Document(paths[0])
    composer = Composer(master)
    for p in paths[1:]:
        master.add_page_break() if False else None
        doc = Document(p)
        # start each chapter on a new page without disturbing its own styles
        first = doc.paragraphs[0]
        run = first.insert_paragraph_before().add_run()
        run.add_break(WD_BREAK.PAGE)
        composer.append(doc)
    composer.save(OUT)

    got_imgs, got_tables = counts(OUT)
    print(f"\n  combined: images {got_imgs} (want {want_imgs}), "
          f"tables {got_tables} (want {want_tables})")

    d = Document(OUT)
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                text += "\n" + c.text

    missing = [h for h in ("Chapter 1: Introduction", "Chapter 2: Literature Review",
                           "Chapter 3: Methodology",
                           "Chapter 4: System Requirement Specification",
                           "Chapter 5: Implementation", "Chapter 6: Testing and Evaluation",
                           "Chapter 7: Concluding Remarks", "References", "Appendices")
               if h not in text]

    problems = []
    if got_imgs != want_imgs:
        problems.append(f"image count {got_imgs} != {want_imgs}")
    if got_tables != want_tables:
        problems.append(f"table count {got_tables} != {want_tables}")
    if missing:
        problems.append(f"headings missing: {missing}")
    if "(VERIFY" in text:
        problems.append("a (VERIFY) placeholder survives")

    for bad, allowed in SUPERSEDED.items():
        for m in re.finditer(re.escape(bad), text):
            ctx = text[max(0, m.start() - 420): m.end() + 420]
            if not any(a in ctx for a in allowed):
                problems.append(f"{bad!r} appears without a supersession note")
                break

    print("\n  verification:")
    if problems:
        for p in problems:
            print(f"    FAIL  {p}")
        raise SystemExit(1)
    print("    all checks passed")
    print(f"\n  wrote {OUT}  ({os.path.getsize(OUT):,} bytes)")


if __name__ == "__main__":
    main()
